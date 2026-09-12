# -*- coding: utf-8 -*-
"""用最优可实施策略D重写 result2.xlsx（计划购电量/充放电量/紧急购电量 三表）"""
import sys, io, json
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")
import numpy as np
import pandas as pd
import openpyxl
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent))
from q2_improve import load_inputs, merge_events, event_label, EMERG_MULT, TOL

BASE = Path(r"C:\Users\fzz17\WorkBuddy\2026-09-10-18-17-23")

price, net_fc, dates, load_act, pv_act, net_act = load_inputs()
z = np.load(BASE / r"results\q2_strategy_D.npz")
B, C, D, EM, SOC0, SOC1 = z["B"], z["C"], z["D"], z["EM"], z["SOC0"], z["SOC1"]
eval_idx = np.where(dates >= pd.Timestamp("2025-02-01"))[0]
plan_cost_day = (B * price).sum(axis=1)

wb = openpyxl.load_workbook(BASE / r"data\附件5\result2.xlsx")

ws = wb["计划购电量"]
assert ws.max_row - 1 == len(eval_idx)
for i, day in enumerate(eval_idx):
    for t in range(144):
        ws.cell(row=i + 2, column=t + 2, value=round(float(B[day, t]), 4))
    ws.cell(row=i + 2, column=146, value=round(float(B[day].sum()), 4))
    ws.cell(row=i + 2, column=147, value=round(float(plan_cost_day[day]), 4))

ws2 = wb["充放电量"]
ws2.delete_rows(2, ws2.max_row - 1)
for day in eval_idx:
    for blk in range(6):
        s, e_ = blk * 24, (blk + 1) * 24
        row = [pd.Timestamp(dates[day]).to_pydatetime() if blk == 0 else None,
               f"{blk*4}:00-{(blk+1)*4}:00",
               round(float(C[day, s:e_].sum()), 4),
               round(float(D[day, s:e_].sum()), 4), None, None]
        if blk == 0:
            row[4] = "0:00"; row[5] = round(float(SOC0[day]), 4)
        elif blk == 1:
            row[4] = "24:00"; row[5] = round(float(SOC1[day]), 4)
        ws2.append(row)

ws3 = wb["紧急购电量"]
ws3.delete_rows(2, ws3.max_row - 1)
for day in eval_idx:
    events = merge_events(EM[day])
    if not events:
        ws3.append([pd.Timestamp(dates[day]).to_pydatetime(), None, None])
    else:
        for j, (s, e_, tot) in enumerate(events):
            ws3.append([pd.Timestamp(dates[day]).to_pydatetime() if j == 0 else None,
                        event_label(s, e_), round(tot, 4)])

out = BASE / r"results\result2.xlsx"
wb.save(out)
print(f"已用策略D重写 {out}")

# 指定日期（D策略，论文表3）
print("\n=== 指定日期紧急购电（D策略） ===")
for ds in ("2025-03-20", "2025-06-21", "2025-09-23", "2025-12-21"):
    day = int(np.where(dates == pd.Timestamp(ds))[0][0])
    evs = merge_events(EM[day])
    txt = "; ".join(f"{event_label(s,e_)}: {tot:.2f} kWh" for s, e_, tot in evs) or "无"
    print(f"  {ds}: {txt}")

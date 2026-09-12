# -*- coding: utf-8 -*-
"""
2026 CUMCM C题 问题1：确定性单日储能经济调度 LP（唯一复现入口）

数据口径（经附件1/2/3/4 交叉核验）：
  - 附件1 时间标记 = 10 分钟区间末端：行 '0:10' 对应区间 [0:00,0:10)，行 '0:00+1' 对应 [23:50,24:00)
  - 功率(kW) 为区间平均功率，电量(kWh) = 功率 × 1/6
  - result1.xlsx 模板与附件1按行位置一一对应

模型（《题目分析报告.md》第5节，确定性储能经济调度族）：
  min  sum_t p_t * g_t
  s.t. E_{t+1} = E_t + eta_c*c_t - d_t/eta_d        t=0..143
       g_t + V_t + d_t = L_t + c_t + w_t            t=0..143
       1200 <= E_t <= 10800                          t=0..144
       0 <= c_t, d_t <= 5000/6                       t=0..143
       g_t, w_t >= 0
       E_0 = E_144 = 6000  （日循环 + 附录1初始电量）
  eta_c = eta_d = 0.9（单向效率，附录1）
  价格非负且允许供能余量 w_t => LP 最优解不会同时充放电（可约化），无需整数变量。
"""
import sys
import io
import json
import hashlib
from pathlib import Path

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")

import numpy as np
import pandas as pd
from scipy.optimize import linprog
from scipy.sparse import lil_matrix, csr_matrix
import openpyxl

BASE = Path(r"C:\Users\fzz17\WorkBuddy\2026-09-10-18-17-23")
DT = 1.0 / 6.0
T = 144
ETA_C = ETA_D = 0.9
E_MIN, E_MAX = 1200.0, 10800.0
P_MAX = 5000.0 * DT      # 833.3333 kWh / 10min
E0 = 6000.0
TOL = 1e-6


def sha256(path):
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def load_attachment1():
    """返回 price[T](元/kWh), load_kw[T], pv_kw[T]；行号 i 对应区间 [i*10min,(i+1)*10min)"""
    df = pd.read_excel(BASE / r"data\附件1.xlsx", sheet_name=0, header=0)
    assert len(df) == 144, f"附件1数据行数异常: {len(df)}"
    price = df.iloc[:, 1].to_numpy(float)
    load_kw = df.iloc[:, 2].to_numpy(float)
    pv_kw = df.iloc[:, 3].to_numpy(float)
    marks = df.iloc[:, 0].tolist()
    assert np.all(price >= 0), "存在负电价，需核查交易边界"
    assert np.all(load_kw >= 0) and np.all(pv_kw >= 0)
    return price, load_kw, pv_kw, marks


def solve_q1(price, L, V):
    """L, V: kWh/区间。返回最优解字典。"""
    # 变量布局: g[0:T], c[0:T], d[0:T], w[0:T], E[0:T+1]
    n = 4 * T + (T + 1)
    ig, ic, id_, iw, iE = 0, T, 2 * T, 3 * T, 4 * T

    cobj = np.zeros(n)
    cobj[ig:ig + T] = price

    # 等式：SOC递推 T 行 + 能量平衡 T 行 + E_0=6000 + E_144=6000
    Aeq = lil_matrix((2 * T + 2, n))
    beq = np.zeros(2 * T + 2)
    for t in range(T):
        r = t
        Aeq[r, iE + t + 1] = 1.0
        Aeq[r, iE + t] = -1.0
        Aeq[r, ic + t] = -ETA_C
        Aeq[r, id_ + t] = 1.0 / ETA_D
        r2 = T + t
        Aeq[r2, ig + t] = 1.0
        Aeq[r2, id_ + t] = 1.0
        Aeq[r2, ic + t] = -1.0
        Aeq[r2, iw + t] = -1.0
        beq[r2] = L[t] - V[t]
    Aeq[2 * T, iE] = 1.0
    beq[2 * T] = E0
    Aeq[2 * T + 1, iE + T] = 1.0
    beq[2 * T + 1] = E0

    bounds = (
        [(0.0, None)] * T +                    # g
        [(0.0, P_MAX)] * T +                   # c
        [(0.0, P_MAX)] * T +                   # d
        [(0.0, None)] * T +                    # w
        [(E_MIN, E_MAX)] * (T + 1)             # E
    )

    res = linprog(cobj, A_eq=csr_matrix(Aeq), b_eq=beq, bounds=bounds, method="highs")
    assert res.status == 0, f"LP未收敛: status={res.status} {res.message}"
    x = res.x
    return {
        "g": x[ig:ig + T], "c": x[ic:ic + T], "d": x[id_:id_ + T],
        "w": x[iw:iw + T], "E": x[iE:iE + T + 1], "cost": float(res.fun),
        "solver_msg": res.message,
    }


def validate(sol, price, L, V):
    """独立复核：能量守恒、SOC边界、日循环、功率上限、互斥、费用复算"""
    g, c, d, w, E = sol["g"], sol["c"], sol["d"], sol["w"], sol["E"]
    checks = {}
    bal = g + V + d - L - c - w
    checks["能量平衡最大残差(kWh)"] = float(np.max(np.abs(bal)))
    soc_res = E[1:] - E[:-1] - ETA_C * c + d / ETA_D
    checks["SOC递推最大残差(kWh)"] = float(np.max(np.abs(soc_res)))
    checks["E最小值(kWh)"] = float(E.min())
    checks["E最大值(kWh)"] = float(E.max())
    checks["E0(kWh)"] = float(E[0])
    checks["E144(kWh)"] = float(E[-1])
    checks["充电功率上限越界"] = float(max(0.0, c.max() - P_MAX))
    checks["放电功率上限越界"] = float(max(0.0, d.max() - P_MAX))
    both = np.minimum(c, d)
    checks["同时充放电最大重叠(kWh)"] = float(both.max())
    cost_recalc = float(np.dot(price, g))
    checks["费用复算(元)"] = cost_recalc
    checks["费用与求解器差(元)"] = abs(cost_recalc - sol["cost"])
    checks["购电总量(kWh)"] = float(g.sum())
    checks["充电总量(kWh)"] = float(c.sum())
    checks["放电总量(kWh)"] = float(d.sum())
    checks["余量总量(kWh)"] = float(w.sum())
    checks["负载总量(kWh)"] = float(L.sum())
    checks["光伏总量(kWh)"] = float(V.sum())
    ok = (checks["能量平衡最大残差(kWh)"] < 1e-6 and checks["SOC递推最大残差(kWh)"] < 1e-6
          and checks["E最小值(kWh)"] >= E_MIN - TOL and checks["E最大值(kWh)"] <= E_MAX + TOL
          and abs(E[0] - E0) < TOL and abs(E[-1] - E0) < TOL
          and checks["同时充放电最大重叠(kWh)"] < 1e-6
          and checks["费用与求解器差(元)"] < 1e-4)
    return checks, ok


def mark_to_str(m):
    return m.strftime("%H:%M") if hasattr(m, "strftime") else str(m)


def interval_label(i):
    """区间 [i*10min, (i+1)*10min) 的标签，如 0:00-0:10"""
    a = i * 10
    b = (i + 1) * 10
    def fmt(mins):
        if mins >= 1440:
            return f"{(mins-1440)//60}:{(mins-1440)%60:02d}+1"
        return f"{mins//60}:{mins%60:02d}"
    return f"{fmt(a)}-{fmt(b)}"


def write_result1(sol, checks):
    src = BASE / r"data\附件5\result1.xlsx"
    out = BASE / r"results\result1.xlsx"
    wb = openpyxl.load_workbook(src)

    ws = wb["计划购电量"]
    assert ws.max_row - 1 == 144, f"模板行数异常: {ws.max_row}"
    for i in range(144):
        ws.cell(row=i + 2, column=2, value=round(float(sol["g"][i]), 4))

    ws2 = wb["充放电量"]
    # 6个4小时块: 0:00-4:00 ... 20:00-24:00，每块24个区间
    for b in range(6):
        s, e = b * 24, (b + 1) * 24
        ws2.cell(row=b + 2, column=2, value=round(float(sol["c"][s:e].sum()), 4))
        ws2.cell(row=b + 2, column=3, value=round(float(sol["d"][s:e].sum()), 4))
    ws2.cell(row=2, column=5, value=round(float(sol["E"][0]), 4))    # 0:00 储电量
    ws2.cell(row=3, column=5, value=round(float(sol["E"][144]), 4))  # 24:00 储电量

    wb.save(out)
    print(f"已写出 {out}")
    return out


def main():
    price, load_kw, pv_kw, marks = load_attachment1()
    L = load_kw * DT
    V = pv_kw * DT

    print("=== 数据概要 ===")
    print(f"负载: 日均功率 {load_kw.mean():.1f} kW, 峰值 {load_kw.max():.1f} kW, 日用电量 {L.sum():.1f} kWh")
    print(f"光伏预测: 峰值 {pv_kw.max():.1f} kW, 日发电量 {V.sum():.1f} kWh")
    print(f"电价: min {price.min():.4f}, max {price.max():.4f}, 均值 {price.mean():.4f} 元/kWh")
    pv_pos = np.where(pv_kw > 0)[0]
    if len(pv_pos):
        print(f"光伏非零区间: {interval_label(pv_pos[0])} 起, {interval_label(pv_pos[-1])} 止, 峰值区间 {interval_label(int(np.argmax(pv_kw)))}")

    sol = solve_q1(price, L, V)
    checks, ok = validate(sol, price, L, V)

    print("\n=== 求解与校验 ===")
    for k, v in checks.items():
        print(f"  {k}: {v:.6f}" if isinstance(v, float) else f"  {k}: {v}")
    print(f"  校验总评: {'PASS' if ok else 'FAIL'}")

    out = write_result1(sol, checks)

    # 论文表1：指定区间购电量
    print("\n=== 表1 指定时间段购电量 ===")
    for hh in (10, 12, 14, 16, 18, 20):
        i = hh * 6  # 区间 [hh:00, hh:10)
        print(f"  {interval_label(i)}: {sol['g'][i]:.4f} kWh")
    print(f"  全天购电量: {sol['g'].sum():.4f} kWh, 全天购电费: {checks['费用复算(元)']:.4f} 元")

    print("\n=== 表2 储能充放电（4小时块） ===")
    for b in range(6):
        s, e = b * 24, (b + 1) * 24
        print(f"  {b*4}:00-{(b+1)*4}:00: 充 {sol['c'][s:e].sum():.4f} kWh, 放 {sol['d'][s:e].sum():.4f} kWh")
    print(f"  0:00 储电量 {sol['E'][0]:.4f} kWh, 24:00 储电量 {sol['E'][144]:.4f} kWh")

    # 保存数值结果供绘图/复核
    np.savez(BASE / r"results\q1_solution.npz",
             g=sol["g"], c=sol["c"], d=sol["d"], w=sol["w"], E=sol["E"],
             price=price, load_kw=load_kw, pv_kw=pv_kw, cost=sol["cost"])

    summary = {
        "question": "q1", "solver": "scipy.optimize.linprog(method=highs)",
        "cost_yuan": checks["费用复算(元)"], "purchase_kwh": checks["购电总量(kWh)"],
        "charge_kwh": checks["充电总量(kWh)"], "discharge_kwh": checks["放电总量(kWh)"],
        "surplus_kwh": checks["余量总量(kWh)"],
        "validation_pass": bool(ok), "checks": checks,
        "input_sha256": {"附件1.xlsx": sha256(BASE / r"data\附件1.xlsx"),
                          "result1模板": sha256(BASE / r"data\附件5\result1.xlsx")},
        "output": str(out),
    }
    with open(BASE / r"results\q1_summary.json", "w", encoding="utf-8") as f:
        json.dump(summary, f, ensure_ascii=False, indent=2)
    print("\n已写出 results/q1_summary.json 与 results/q1_solution.npz")
    return 0 if ok else 1


if __name__ == "__main__":
    sys.exit(main())

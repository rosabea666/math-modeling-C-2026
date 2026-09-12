# -*- coding: utf-8 -*-
"""检查附件1(Q1数据)与result1.xlsx模板结构"""
import sys
import io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")
import openpyxl

BASE = r"C:\Users\fzz17\WorkBuddy\2026-09-10-18-17-23\data"

def dump(path, name, max_rows=12):
    wb = openpyxl.load_workbook(path, data_only=True)
    print(f"##### {name} #####")
    for ws in wb.worksheets:
        print(f"-- sheet: {ws.title}  dims={ws.dimensions} max_row={ws.max_row} max_col={ws.max_column}")
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i >= max_rows:
                print("   ...")
                break
            print("  ", row)
        # last rows
        if ws.max_row > max_rows:
            tail = list(ws.iter_rows(min_row=max(1, ws.max_row - 2), values_only=True))
            print("   TAIL:")
            for r in tail:
                print("  ", r)
    print()

dump(BASE + r"\附件1.xlsx", "附件1")
dump(BASE + r"\附件5\result1.xlsx", "result1模板", max_rows=8)

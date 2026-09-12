# -*- coding: utf-8 -*-
"""用主方案 E（V_dow_N）重写 result2.xlsx，并生成论文表1/表2/表3 数据。

与 q2_finalize.py 的区别：
  · 预测器 = 星期偏差校正（forecast_dow_N），逐月 walk-forward 日程
  · 计划层 = make_plan（day-cycle，E_hat_H = E_0）
  · 执行层 = must/max 尽限充放电，实际 SOC 跨日连续
  · 2/1 起评，E = 10800 kWh（公共预热口径）

只写 results/result2.xlsx 与 results/q2_E_tables.json，不改模型。
"""
import sys, io, json
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")
import numpy as np
import pandas as pd
import openpyxl
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent))
from q2_model import (T, MULT, E_FEB1, WARMUP_DAY, E_MIN, E_MAX,  # noqa: E402
                      Policy, load_data, make_plan, execute)
from q2_improve import merge_events, event_label  # noqa: E402

BASE = Path(__file__).resolve().parent.parent

POL = Policy("E", forecast="cquant", W=7, beta=0.70, exec_mode="greedy",
             discharge_policy="must", charge_policy="max")

# ---------- 星期偏差校正预测器（与 scripts/q2_dow_N_joint.py 逐字一致） ----------
def _col_quantile(X, qs):
    """X: (n, T)，qs: (T,)。返回每列在 qs 对应分位上的值，形状 (T,)。"""
    out = np.empty(X.shape[1])
    for j in range(X.shape[1]):
        out[j] = np.quantile(X[:, j], qs[j])
    return out


def dow_index(dates):
    return np.array([pd.Timestamp(d).dayofweek for d in dates])


def m_recent(d, X, W):
    hist = X[max(0, d - W):d]
    if hist.shape[0] < 7:
        if d >= 30:
            return np.median(X[max(0, d - 30):d], axis=0)
        return np.zeros(X.shape[1])
    return np.median(hist, axis=0)


def delta_dow(d, X, dow_idx, n_dow_hist, W_level):
    w_d = dow_idx[d]
    cand = [d_ for d_ in range(max(0, d - 90), d) if dow_idx[d_] == w_d]
    cand = cand[-n_dow_hist:]
    if not cand:
        return np.zeros(X.shape[1])
    deltas = np.zeros((len(cand), X.shape[1]))
    for i, d_ in enumerate(cand):
        m_ = m_recent(d_, X, W_level)
        deltas[i] = X[d_] - m_
    return deltas.mean(axis=0)


def forecast_dow_N(d, D, dow_idx, gamma, W_level, beta, n_dow_hist):
    if d < 30:
        return D["net_ref"].copy()
    m = m_recent(d, D["net"], W_level)
    dlt = delta_dow(d, D["net"], dow_idx, n_dow_hist, W_level)
    base = m + gamma * dlt
    net_hist = D["net"][max(0, d - W_level):d]
    if net_hist.shape[0] < 7:
        return D["net_ref"].copy()
    resid = net_hist - np.median(net_hist, axis=0)
    margin = _col_quantile(resid, np.full(T, beta))
    return base + margin


# ---------- §37 联合 walk-forward 日程 ----------
SCHED = {m: (0.70, 7, 1.0, 4) for m in (2, 3)}
for m in range(4, 13):
    SCHED[m] = (0.75, 7, 1.0, 4)


def main():
    D = load_data()
    dow_idx = dow_index(D["dates"])
    months = np.array([d.astype("datetime64[M]").astype(int) % 12 + 1
                       for d in D["dates"]])
    price, net = D["price"], D["net"]
    nd = len(D["dates"])
    idx = np.where(D["dates"] >= np.datetime64("2025-02-01"))[0]

    B = np.zeros((nd, T)); C = np.zeros((nd, T))
    DD = np.zeros((nd, T)); EM = np.zeros((nd, T)); W_ = np.zeros((nd, T))
    SOC0 = np.zeros(nd); SOC1 = np.zeros(nd)

    E = float(E_FEB1)
    for d in range(WARMUP_DAY, nd):
        beta, W_level, gamma, n_dow = SCHED[months[d]]
        N_hat = forecast_dow_N(d, D, dow_idx, gamma, W_level, beta, n_dow)
        b, c_plan = make_plan(POL, price, N_hat, E)
        c, dd, e, w, Etraj = execute(POL, b, c_plan, N_hat, net[d], price, E)
        B[d], C[d], DD[d], EM[d], W_[d] = b, c, dd, e, w
        SOC0[d], SOC1[d] = Etraj[0], Etraj[-1]
        E = Etraj[-1]

    ev = idx
    plan_cost = float((B[ev] * price).sum()) / 1e4
    em_cost = float(MULT * (EM[ev] * price).sum()) / 1e4
    total = plan_cost + em_cost
    print("=" * 78)
    print("主方案 E（V_dow_N）评价期回放")
    print("=" * 78)
    print(f"  评价期 {len(ev)} 天  2/1 起评 E = {E_FEB1:.0f} kWh")
    print(f"  计划费 {plan_cost:.4f} 万元 | 紧急费 {em_cost:.4f} 万元 | "
          f"总费 {total:.4f} 万元")
    assert abs(total - 1395.6995) < 0.01, f"与 §37 不一致：{total}"

    # ---------- 表1：指定日期计划购电量 ----------
    slots = {'10:00-10:10': 60, '12:00-12:10': 72, '14:00-14:10': 84,
             '16:00-16:10': 96, '18:00-18:10': 108, '20:00-20:10': 120}
    tbl1 = {}
    print("\n=== 表1 指定日期计划购电量（kWh） ===")
    for ds in ("2025-03-20", "2025-06-21", "2025-09-23", "2025-12-21"):
        day = int(np.where(D["dates"] == np.datetime64(ds))[0][0])
        b_slots = {k: float(B[day, t]) for k, t in slots.items()}
        b_total = float(B[day].sum())
        em_kwh = float(EM[day].sum())
        pc = float((B[day] * price).sum())
        ec = float(MULT * (EM[day] * price).sum())
        tbl1[ds] = {"b_slots": b_slots, "计划购电量": b_total,
                    "紧急购电量": em_kwh, "合计购电量": b_total + em_kwh,
                    "计划费_元": pc, "紧急费_元": ec, "全日费用_元": pc + ec}
        print(f"  {ds}: " + "/".join(f"{b_slots[k]:.4f}" for k in slots)
              + f" | 计划 {b_total:.4f} kWh / 紧急 {em_kwh:.4f} kWh "
              + f"/ 合计 {b_total+em_kwh:.4f} kWh | ¥{pc+ec:.4f}")

    # ---------- 表2：指定日期充放电量与日初日末 ----------
    tbl2 = {}
    print("\n=== 表2 指定日期充放电量（kWh） ===")
    for ds in ("2025-03-20", "2025-06-21", "2025-09-23", "2025-12-21"):
        day = int(np.where(D["dates"] == np.datetime64(ds))[0][0])
        blocks = {}
        for blk in range(6):
            s, e_ = blk * 24, (blk + 1) * 24
            blocks[f"{blk*4}:00-{(blk+1)*4}:00"] = {
                "充": float(C[day, s:e_].sum()), "放": float(DD[day, s:e_].sum())}
        tbl2[ds] = {"blocks": blocks, "日初": float(SOC0[day]),
                    "日末": float(SOC1[day])}
        line = " | ".join(f"{k}:{v['充']:.1f}/{v['放']:.1f}"
                          for k, v in blocks.items())
        print(f"  {ds}: {line}")
        print(f"      日初 {SOC0[day]:.4f} / 日末 {SOC1[day]:.4f} kWh")

    # ---------- 表3：指定日期紧急购电（合并连续区间） ----------
    tbl3 = {}
    print("\n=== 表3 指定日期紧急购电（合并连续区间） ===")
    for ds in ("2025-03-20", "2025-06-21", "2025-09-23", "2025-12-21"):
        day = int(np.where(D["dates"] == np.datetime64(ds))[0][0])
        evs = merge_events(EM[day])
        tbl3[ds] = [{"区间": event_label(s, e_), "电量_kWh": float(tot)}
                    for s, e_, tot in evs]
        txt = "; ".join(f"{event_label(s, e_)}: {tot:.2f} kWh"
                        for s, e_, tot in evs) or "无紧急购电"
        print(f"  {ds}（{len(evs)} 个事件，共 {EM[day].sum():.2f} kWh）: {txt}")

    json.dump({"总费_万元": total, "计划费_万元": plan_cost,
               "紧急费_万元": em_cost, "表1": tbl1, "表2": tbl2, "表3": tbl3},
              open(BASE / r"results\q2_E_tables.json", "w", encoding="utf-8"),
              ensure_ascii=False, indent=1)

    # ---------- 写 result2.xlsx ----------
    wb = openpyxl.load_workbook(BASE / r"data\附件5\result2.xlsx")
    plan_cost_day = (B * price).sum(axis=1)

    ws = wb["计划购电量"]
    assert ws.max_row - 1 == len(ev)
    for i, day in enumerate(ev):
        for t in range(T):
            ws.cell(row=i + 2, column=t + 2, value=round(float(B[day, t]), 4))
        ws.cell(row=i + 2, column=146, value=round(float(B[day].sum()), 4))
        ws.cell(row=i + 2, column=147, value=round(float(plan_cost_day[day]), 4))

    ws2 = wb["充放电量"]
    ws2.delete_rows(2, ws2.max_row - 1)
    for day in ev:
        for blk in range(6):
            s, e_ = blk * 24, (blk + 1) * 24
            row = [pd.Timestamp(D["dates"][day]).to_pydatetime() if blk == 0 else None,
                   f"{blk*4}:00-{(blk+1)*4}:00",
                   round(float(C[day, s:e_].sum()), 4),
                   round(float(DD[day, s:e_].sum()), 4), None, None]
            if blk == 0:
                row[4] = "0:00"; row[5] = round(float(SOC0[day]), 4)
            elif blk == 1:
                row[4] = "24:00"; row[5] = round(float(SOC1[day]), 4)
            ws2.append(row)

    ws3 = wb["紧急购电量"]
    ws3.delete_rows(2, ws3.max_row - 1)
    for day in ev:
        evs = merge_events(EM[day])
        if not evs:
            ws3.append([pd.Timestamp(D["dates"][day]).to_pydatetime(), None, None])
        else:
            for j, (s, e_, tot) in enumerate(evs):
                ws3.append([pd.Timestamp(D["dates"][day]).to_pydatetime()
                            if j == 0 else None,
                            event_label(s, e_), round(tot, 4)])

    out = BASE / r"results\result2.xlsx"
    wb.save(out)
    print(f"\n已用主方案 E（V_dow_N）重写 {out}")


if __name__ == "__main__":
    main()

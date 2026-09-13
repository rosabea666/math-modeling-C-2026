# -*- coding: utf-8 -*-
"""问题三最终链路：跑 B1 主方案，导出论文表 1/2/3 与 result3.xlsx。

与 q2_finalize_E.py 同体例：只写 results/ 下的产物，不改模型。
输出：
  results/q3_tables.json   四个指定日期的表 1/表 2/表 3 数据 + 2×2 矩阵
  results/result3.xlsx     按 data/附件5/result3.xlsx 模板填写

口径（与论文 §7 一致，且 2×2 四项已与论文数逐位吻合）：
  计划购电量 = 0:00 合同 b_plan；调整购电量 = 最终生效合同 b'
  结算 J = Σ p·b' + ½Σ p|b'−b_plan| + 5Σ p·e
"""
import sys
import io
import json
from pathlib import Path

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", line_buffering=True)

import numpy as np
import pandas as pd
import openpyxl

sys.path.insert(0, str(Path(__file__).parent))
import q3_core as Q                                       # noqa: E402
from q2_improve import merge_events, event_label          # noqa: E402

BASE = Path(__file__).resolve().parent.parent
DATES = ("2025-03-20", "2025-06-21", "2025-09-23", "2025-12-21")
SLOTS = {"10:00-10:10": 60, "12:00-12:10": 72, "14:00-14:10": 84,
         "16:00-16:10": 96, "18:00-18:10": 108, "20:00-20:10": 120}
BLOCK_NAMES = [f"{h * 4}:00-{h * 4 + 4}:00" for h in range(6)]


def main():
    D = Q.load_data()
    price = D["price"]
    idx = np.where(D["dates"] >= np.datetime64("2025-02-01"))[0]

    print("=" * 74)
    print("问题三全链路复算")
    print("=" * 74)
    R = {}
    for tag, kw in [("A0", dict(use_channel=False, allow_adjust=False)),
                    ("A0K", dict(use_channel=True, allow_adjust=False)),
                    ("B0", dict(use_channel=False, allow_adjust=True)),
                    ("B1", dict(use_channel=True, allow_adjust=True))]:
        R[tag] = Q.replay(D, label=tag, **kw)
        print("  审计", Q.audit(R[tag], D, tag))
    B1 = R["B1"]

    # ---------- 表1：指定日期计划购电量 ----------
    tbl1 = {}
    print("\n=== 表1 指定日期计划购电量（kWh）===")
    for ds in DATES:
        day = int(np.where(D["dates"] == np.datetime64(ds))[0][0])
        b = B1["B"][day]
        slots = {k: float(b[t]) for k, t in SLOTS.items()}
        em_kwh = float(B1["EM"][day].sum())
        bf = B1["BF"][day]
        tbl1[ds] = {"b_slots": slots,
                    "bf_slots": {k: float(bf[t]) for k, t in SLOTS.items()},
                    "计划购电量": float(b.sum()),
                    "调整后购电量": float(bf.sum()),
                    "紧急购电量": em_kwh,
                    "合计购电量": float(bf.sum()) + em_kwh,
                    "计划购电费_元": float(np.sum(price * b)),
                    "调整附加费_元": float(0.5 * np.sum(price * np.abs(bf - b))),
                    "紧急购电费_元": float(Q.MULT * np.sum(price * B1["EM"][day]))}
        print(f"  {ds}: " + "/".join(f"{slots[k]:.2f}" for k in SLOTS)
              + f" | 计划 {b.sum():.1f} kWh / 紧急 {em_kwh:.1f} kWh")

    # ---------- 表2：指定日期充放电量与日初日末 ----------
    tbl2 = {}
    print("\n=== 表2 指定日期充放电量（kWh）===")
    for ds in DATES:
        day = int(np.where(D["dates"] == np.datetime64(ds))[0][0])
        blocks = {}
        for blk in range(6):
            s, e_ = blk * 24, (blk + 1) * 24
            blocks[BLOCK_NAMES[blk]] = {
                "充": float(B1["C"][day, s:e_].sum()),
                "放": float(B1["D"][day, s:e_].sum())}
        tbl2[ds] = {"blocks": blocks, "日初": float(B1["SOC0"][day]),
                    "日末": float(B1["SOC1"][day])}
        print(f"  {ds}: " + " | ".join(f"{k} {v['充']:.1f}/{v['放']:.1f}"
                                       for k, v in blocks.items()))
        print(f"      日初 {B1['SOC0'][day]:.2f} / 日末 {B1['SOC1'][day]:.2f} kWh")

    # ---------- 表3：指定日期紧急购电 ----------
    tbl3 = {}
    print("\n=== 表3 指定日期紧急购电（合并连续区间）===")
    for ds in DATES:
        day = int(np.where(D["dates"] == np.datetime64(ds))[0][0])
        evs = merge_events(B1["EM"][day])
        tbl3[ds] = [{"区间": event_label(s, e_), "电量_kWh": float(tot)}
                    for s, e_, tot in evs]
        txt = "; ".join(f"{event_label(s, e_)} {tot:.1f} kWh" for s, e_, tot in evs)
        print(f"  {ds}（{len(evs)} 个事件，共 {B1['EM'][day].sum():.1f} kWh）: "
              f"{txt or '无紧急购电'}")

    matrix = {k: R[k]["costs"]["总费_万元"] for k in R}
    out = {"总费_万元": matrix["B1"], "矩阵": matrix,
           "分项_万元": {k: {kk: vv for kk, vv in R[k]["costs"].items()}
                      for k in R},
           "表1": tbl1, "表2": tbl2, "表3": tbl3}
    js = BASE / "results" / "q3_tables.json"
    json.dump(out, open(js, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    print(f"\n已写 {js}")

    # ---------- 写 result3.xlsx ----------
    wb = openpyxl.load_workbook(BASE / "data" / "附件5" / "result3.xlsx")
    ws = wb["计划购电量"]
    ws2 = wb["调整购电量"]
    assert ws.max_row - 1 == len(idx) == ws2.max_row - 1, "模板行数与评价期不符"
    for i, day in enumerate(idx):
        bp = B1["B"][day]
        bf = B1["BF"][day]
        for t in range(Q.T):
            ws.cell(row=i + 2, column=t + 2, value=round(float(bp[t]), 4))
            ws2.cell(row=i + 2, column=t + 2, value=round(float(bf[t]), 4))
        ws.cell(row=i + 2, column=146, value=round(float(bp.sum()), 4))
        ws.cell(row=i + 2, column=147,
                value=round(float(np.sum(price * bp)), 4))
        ws2.cell(row=i + 2, column=146, value=round(float(bf.sum()), 4))
        ws2.cell(row=i + 2, column=147,
                 value=round(float(np.sum(price * bf)
                                   + 0.5 * np.sum(price * np.abs(bf - bp))), 4))

    ws3 = wb["充放电量"]
    ws3.delete_rows(2, ws3.max_row - 1)
    for day in idx:
        for blk in range(6):
            s, e_ = blk * 24, (blk + 1) * 24
            row = [pd.Timestamp(D["dates"][day]).to_pydatetime() if blk == 0 else None,
                   BLOCK_NAMES[blk],
                   round(float(B1["C"][day, s:e_].sum()), 4),
                   round(float(B1["D"][day, s:e_].sum()), 4), None, None]
            if blk == 0:
                row[4] = "0:00"; row[5] = round(float(B1["SOC0"][day]), 4)
            elif blk == 1:
                row[4] = "24:00"; row[5] = round(float(B1["SOC1"][day]), 4)
            ws3.append(row)

    ws4 = wb["紧急购电量"]
    ws4.delete_rows(2, ws4.max_row - 1)
    for day in idx:
        evs = merge_events(B1["EM"][day])
        if not evs:
            ws4.append([pd.Timestamp(D["dates"][day]).to_pydatetime(), None, None])
        else:
            for j, (s, e_, tot) in enumerate(evs):
                ws4.append([pd.Timestamp(D["dates"][day]).to_pydatetime()
                            if j == 0 else None,
                            event_label(s, e_), round(float(tot), 4)])

    out_x = BASE / "results" / "result3.xlsx"
    wb.save(out_x)
    print(f"已写 {out_x}")


if __name__ == "__main__":
    main()
